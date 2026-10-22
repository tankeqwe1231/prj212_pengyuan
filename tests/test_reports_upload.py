import ast
from datetime import datetime
import io
import json
import os
from pathlib import Path
import re
import signal
import subprocess
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch
from urllib import parse
import uuid

import fitz
from flask import (Flask, Response, flash, g, jsonify, make_response,
                   redirect, request, send_file, url_for)
from flask_appbuilder import expose
from flask_appbuilder.security.decorators import has_access
from flask_babel import Babel
from flask_wtf import CSRFProtect
from flask_wtf.csrf import generate_csrf
import sqlalchemy as sqla


MODULE_PATH = Path(__file__).parents[1] / 'superset' / 'views' / 'core_custom.py'


def load_reports_view(root):
    # Extract methods without importing Superset, deployment config or services.
    tree = ast.parse(MODULE_PATH.read_text(encoding='utf-8'))
    view = next(node for node in tree.body
                if isinstance(node, ast.ClassDef) and node.name == 'ReportsDataManage')
    methods = [node for node in view.body if isinstance(node, ast.FunctionDef)]
    namespace = {
        '__file__': str(root / 'superset' / 'views' / 'core_custom.py'),
        'datetime': datetime, 'fitz': fitz, 'json': json, 'os': os, 'sys': sys,
        're': re, 'uuid': uuid, 'parse': parse, 'sqla': sqla,
        'flash': flash, 'g': g, 'jsonify': jsonify, 'request': request,
        'redirect': redirect, 'url_for': url_for, 'send_file': send_file,
        'make_response': make_response, 'expose': expose,
        'has_access': has_access, 'log_this': lambda method: method,
        'logging': Mock(), 'traceback': Mock(format_exc=lambda: 'isolated failure'),
        'db': SimpleNamespace(session=Mock(), engine=Mock()),
        'config': {
            'OSS_BUCKET_FILES': 'test-bucket', 'OSS_ENDPOINT': 'example.invalid',
            'OSS_FILE_CUSTOM_DOMAIN': 'example.invalid',
        },
        'AcsClient': Mock(), 'oss2': Mock(),
        'generate_oss_preview_url': Mock(return_value='https://example.invalid/file'),
        'json_error_response': lambda message, status=500: Response(
            json.dumps({'error': message}), status=status, mimetype='application/json'),
    }
    module = ast.Module(body=methods, type_ignores=[])
    ast.fix_missing_locations(module)
    exec(compile(module, str(MODULE_PATH), 'exec'), namespace)
    instance = type('ReportsDataManage', (), {
        method.name: namespace[method.name] for method in methods})()
    instance.base_permissions = ['can_index', 'can_get_report_file_content',
                                 'can_download_report_file']
    instance.class_permission_name = 'ReportsDataManage'
    instance.method_permission_name = None
    instance.appbuilder = SimpleNamespace(sm=SimpleNamespace(
        has_access=lambda permission, view_name: g.user.allowed,
        auth_view=type('TestAuthView', (), {})(),
    ))
    return instance, namespace


def pdf_bytes():
    with fitz.open() as document:
        first = document.new_page(width=600, height=900)
        first.draw_rect(first.rect, color=(1, 0, 0), fill=(1, 0, 0))
        second = document.new_page(width=600, height=900)
        second.draw_rect(second.rect, color=(0, 0, 1), fill=(0, 0, 1))
        document.set_metadata({'creator': 'Microsoft Word'})
        return document.tobytes()


def pptx_bytes(count=2, hide_last=False):
    from pptx import Presentation

    presentation = Presentation()
    for index in range(count):
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = 'Report slide {}'.format(index + 1)
        if hide_last and index == count - 1:
            slide._element.set('show', '0')
    stream = io.BytesIO()
    presentation.save(stream)
    return stream.getvalue()


class ReportsUploadTest(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.upload_dir = self.root / 'superset' / 'views' / 'reports_upload'
        self.upload_dir.mkdir(parents=True)
        self.view, self.namespace = load_reports_view(self.root)
        self.view._download_from_oss = Mock(side_effect=AssertionError('OSS called'))
        self.view._get_oss_upload_credentials = Mock(side_effect=AssertionError('STS called'))
        self.app = Flask(__name__)
        self.app.config.update(TESTING=True, SECRET_KEY='isolated-test-only')
        Babel(self.app)
        CSRFProtect(self.app)
        self.allowed = True

        @self.app.before_request
        def test_user():
            g.user = SimpleNamespace(id=7, allowed=self.allowed)

        self.app.add_url_rule('/csrf', view_func=lambda: jsonify(token=generate_csrf()))
        self.app.add_url_rule('/login', endpoint='TestAuthView.login', view_func=lambda: '')
        self.app.add_url_rule('/list', view_func=self.view.index, methods=['POST'])
        self.app.add_url_rule('/content', view_func=self.view.get_report_file_content)
        self.app.add_url_rule('/download', view_func=self.view.download_report_file)
        self.client = self.app.test_client()
        token = self.client.get('/csrf').get_json()['token']
        self.headers = {'X-CSRFToken': token}

    def tearDown(self):
        self.view._download_from_oss.assert_not_called()
        self.view._get_oss_upload_credentials.assert_not_called()
        self.namespace['AcsClient'].assert_not_called()
        self.namespace['oss2'].Auth.assert_not_called()

    def upload(self, content, filename):
        response = self.client.post('/list', data={
            'action': 'upload_report_file', 'file': (io.BytesIO(content), filename),
        }, headers=self.headers)
        response.request.environ['wsgi.input'].close()
        return response

    def submit(self, action='add_report', **overrides):
        data = {
            'action': action, 'report_title': 'Test report',
            'report_file_url': 'a' * 32, 'report_file_suffix': 'pdf',
            'report_file_origin_name': '报告.pdf',
        }
        data.update(overrides)
        return self.client.post('/list', json=data, headers=self.headers)

    def saved_pdf(self):
        path = self.upload_dir / ('a' * 32 + '.pdf')
        path.write_bytes(pdf_bytes())
        return path

    def test_pdf_upload_stays_local_and_returns_pdf_metadata(self):
        content = pdf_bytes()
        response = self.upload(content, '季度报告.PDF')
        self.assertEqual(200, response.status_code)
        data = response.get_json()
        self.assertEqual('季度报告.pdf', data['file_origin_name'])
        self.assertEqual('pdf', data['file_suffix'])
        self.assertEqual(content, (self.upload_dir / (data['file_uuid'] + '.pdf')).read_bytes())
        self.namespace['db'].session.execute.assert_not_called()

    def test_pptx_upload_converts_and_preserves_source(self):
        def convert(source, target):
            self.assertEqual(b'pptx source', Path(source).read_bytes())
            Path(target).write_bytes(pdf_bytes())

        self.view._convert_pptx_to_pdf = Mock(side_effect=convert)
        response = self.upload(b'pptx source', '../季度报告.PPTX')
        self.assertEqual(200, response.status_code)
        data = response.get_json()
        self.assertEqual('季度报告.pdf', data['file_origin_name'])
        self.assertEqual('pdf', data['file_suffix'])
        self.assertEqual(b'pptx source', (self.upload_dir / (data['file_uuid'] + '.pptx')).read_bytes())
        self.assertTrue((self.upload_dir / (data['file_uuid'] + '.pdf')).is_file())

    def test_upload_requires_file_and_rejects_video_empty_and_invalid_pdf(self):
        missing = self.client.post('/list', data={'action': 'upload_report_file'}, headers=self.headers)
        self.assertEqual(400, missing.status_code)
        for content, filename in ((b'video', 'report.mp4'), (b'', 'report.pptx'),
                                  (b'not pdf', 'report.pdf')):
            with self.subTest(filename=filename):
                self.assertEqual(400, self.upload(content, filename).status_code)
        self.assertEqual([], list(self.upload_dir.iterdir()))

    def test_pdf_size_limit_is_enforced_on_server(self):
        response = self.upload(b'x' * (20 * 1024 * 1024 + 1), 'large.pdf')
        self.assertEqual(400, response.status_code)
        self.assertEqual([], list(self.upload_dir.iterdir()))

    def test_encrypted_pdf_is_rejected(self):
        with fitz.open(stream=pdf_bytes(), filetype='pdf') as document:
            encrypted = document.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256,
                                         owner_pw='test', user_pw='test')
        self.assertEqual(400, self.upload(encrypted, 'encrypted.pdf').status_code)

    def test_conversion_failure_retains_original_and_does_not_save_report(self):
        self.view._convert_pptx_to_pdf = Mock(side_effect=RuntimeError('conversion failed'))
        response = self.upload(b'pptx source', 'report.pptx')
        self.assertEqual(500, response.status_code)
        self.assertEqual(1, len(list(self.upload_dir.glob('*.pptx'))))
        self.namespace['db'].session.execute.assert_not_called()

    def test_add_uses_local_pdf_null_oss_key_and_full_first_page_cover(self):
        path = self.saved_pdf()
        response = self.submit()
        self.assertEqual(200, response.status_code)
        params = self.namespace['db'].session.execute.call_args.args[1]
        self.assertEqual(str(path), params['report_file_url'])
        self.assertEqual('pdf', params['report_file_suffix'])
        self.assertIsNone(params['oss_key'])
        self.namespace['db'].session.commit.assert_called_once()
        cover = self.root / 'superset' / 'assets' / 'images' / 'reports_covers' / Path(params['cover_img_url']).name
        pixmap = fitz.Pixmap(str(cover))
        self.assertEqual((1080, 1620), (pixmap.width, pixmap.height))
        self.assertEqual((255, 0, 0), pixmap.pixel(100, 1500))

    def test_custom_cover_is_preserved(self):
        self.saved_pdf()
        self.view._generate_cover_from_pdf = Mock()
        self.assertEqual(200, self.submit(cover_img_url='/custom.png').status_code)
        self.view._generate_cover_from_pdf.assert_not_called()
        params = self.namespace['db'].session.execute.call_args.args[1]
        self.assertEqual('/custom.png', params['cover_img_url'])

    def test_update_replacement_clears_oss_and_uses_pdf(self):
        path = self.saved_pdf()
        self.assertEqual(200, self.submit(action='update_report', id=4).status_code)
        params = self.namespace['db'].session.execute.call_args.args[1]
        self.assertEqual(str(path), params['report_file_url'])
        self.assertEqual('pdf', params['report_file_suffix'])
        self.assertIsNone(params['oss_key'])
        self.assertTrue(params['cover_img_url'].endswith('.png'))

    def test_metadata_only_update_keeps_existing_file_and_cover(self):
        self.assertEqual(200, self.submit(action='update_report', id=4, report_file_url=None).status_code)
        sql, params = self.namespace['db'].session.execute.call_args.args
        for field in ('report_file_url', 'report_file_suffix', 'oss_key', 'cover_img_url'):
            self.assertNotIn(field, params)
            self.assertNotIn(field, str(sql))

    def test_invalid_or_missing_file_never_reaches_database(self):
        for file_id in ('../escape', 'C:\\outside\\secret', 'b' * 32, None):
            with self.subTest(file_id=file_id):
                self.assertNotEqual(200, self.submit(report_file_url=file_id).status_code)
        self.namespace['db'].session.execute.assert_not_called()

    def test_database_failure_rolls_back(self):
        self.saved_pdf()
        self.namespace['db'].session.execute.side_effect = RuntimeError('database unavailable')
        self.assertEqual(500, self.submit().status_code)
        self.namespace['db'].session.rollback.assert_called_once()
        self.namespace['db'].session.commit.assert_not_called()

    def test_missing_csrf_and_denied_permission_do_not_upload(self):
        response = self.client.post('/list', data={
            'action': 'upload_report_file', 'file': (io.BytesIO(pdf_bytes()), 'report.pdf'),
        })
        self.assertEqual(400, response.status_code)
        # FAB's real decorator delegates authentication/authorization to sm.
        self.allowed = False
        self.assertEqual(302, self.upload(pdf_bytes(), 'report.pdf').status_code)
        self.assertEqual([], list(self.upload_dir.iterdir()))

    def test_local_preview_and_download_return_pdf_without_oss_redirect(self):
        path = self.saved_pdf()
        self.namespace['db'].engine.execute.return_value.fetchone.return_value = {
            'report_file_url': str(path), 'report_file_origin_name': '季度报告.pdf',
            'report_file_suffix': 'pdf', 'oss_key': None,
        }
        for route, disposition in (('/content', 'inline;'), ('/download', 'attachment;')):
            response = self.client.get(route + '?report_id=4')
            self.assertEqual(200, response.status_code)
            self.assertTrue(response.headers['Content-Disposition'].startswith(disposition))
            self.assertIn('%E5%AD%A3%E5%BA%A6%E6%8A%A5%E5%91%8A.pdf', response.headers['Content-Disposition'])
            self.assertEqual(path.read_bytes(), response.data)
            response.close()
        self.namespace['generate_oss_preview_url'].assert_not_called()

    def test_existing_oss_report_preview_still_works(self):
        self.namespace['db'].engine.execute.return_value.fetchone.return_value = {
            'report_file_url': '', 'report_file_origin_name': 'old.pdf',
            'report_file_suffix': 'pdf', 'oss_key': 'old.pdf',
        }
        response = self.client.get('/content?report_id=4')
        self.assertEqual(302, response.status_code)
        self.assertEqual('pdfs/old.pdf', self.namespace['generate_oss_preview_url'].call_args.kwargs['object_key'])

    def test_real_libreoffice_conversion_includes_hidden_slide_without_spire(self):
        try:
            self.view._get_libreoffice_executable()
        except RuntimeError as error:
            self.skipTest(str(error))
        content = pptx_bytes(count=4, hide_last=True)
        with patch.dict(sys.modules, {'spire': None, 'spire.presentation': None}):
            response = self.upload(content, '四页报告 with hidden slide.pptx')
        self.assertEqual(200, response.status_code, response.get_json())
        file_uuid = response.get_json()['file_uuid']
        with fitz.open(self.upload_dir / (file_uuid + '.pdf')) as document:
            self.assertEqual(4, document.page_count)
            for index, page in enumerate(document):
                self.assertIn('Report slide {}'.format(index + 1), page.get_text())
        self.assertEqual(content, (self.upload_dir / (file_uuid + '.pptx')).read_bytes())
        self.assertEqual([], list(self.upload_dir.glob('reports-libreoffice-*')))


class ReportsConversionTest(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.view, self.namespace = load_reports_view(self.root)
        self.view._get_libreoffice_executable = Mock(return_value='test-soffice')
        self.source = self.root / '季度报告 with spaces.PPTX'
        self.target = self.root / 'renamed.pdf'
        self.original = pptx_bytes()
        self.source.write_bytes(self.original)
        popen_patch = patch('subprocess.Popen')
        self.popen = popen_patch.start()
        self.addCleanup(popen_patch.stop)
        self.process = self.popen.return_value.__enter__.return_value
        self.process.pid = 43210
        self.process.returncode = 0
        self.process.communicate.side_effect = self.write_converted_pdf
        self.output = pdf_bytes()

    def write_converted_pdf(self, **kwargs):
        command = self.popen.call_args.args[0]
        output_dir = Path(command[command.index('--outdir') + 1])
        (output_dir / (self.source.stem + '.pdf')).write_bytes(self.output)
        return b'', b''

    def convert(self):
        self.view._convert_pptx_to_pdf(str(self.source), str(self.target))

    def assert_clean_failure(self):
        self.assertFalse(self.target.exists())
        self.assertEqual(self.original, self.source.read_bytes())
        self.assertEqual([], list(self.root.glob('reports-libreoffice-*')))

    def test_windows_finds_default_install_without_path_entry(self):
        program_files = str(self.root / 'Program Files')
        executable = os.path.join(program_files, 'LibreOffice', 'program', 'soffice.com')
        with patch.dict(self.namespace, {'sys': SimpleNamespace(platform='win32')}), \
                patch.dict(os.environ, {'ProgramFiles': program_files}, clear=True), \
                patch('shutil.which', return_value=None), \
                patch('os.path.isfile', side_effect=lambda path: path == executable):
            self.assertEqual(executable, type(self.view)._get_libreoffice_executable(self.view))

    def test_linux_finds_libreoffice_or_soffice_on_path(self):
        for name in ('libreoffice', 'soffice'):
            with self.subTest(name=name), \
                    patch.dict(self.namespace, {'sys': SimpleNamespace(platform='linux')}), \
                    patch('shutil.which', side_effect=lambda candidate: '/usr/bin/' + name if candidate == name else None):
                self.assertEqual('/usr/bin/' + name,
                                 type(self.view)._get_libreoffice_executable(self.view))

    def test_missing_installation_has_clear_error(self):
        with patch.dict(self.namespace, {'sys': SimpleNamespace(platform='linux')}), \
                patch('shutil.which', return_value=None):
            with self.assertRaisesRegex(RuntimeError, '未找到 LibreOffice'):
                type(self.view)._get_libreoffice_executable(self.view)

    def test_conversion_preserves_source_and_supports_different_pdf_name(self):
        self.convert()
        self.assertEqual(self.output, self.target.read_bytes())
        self.assertEqual(self.original, self.source.read_bytes())
        self.assertEqual([], list(self.root.glob('reports-libreoffice-*')))
        command = self.popen.call_args.args[0]
        self.assertEqual(str(self.source), command[-1])
        self.assertIn('pdf:impress_pdf_Export', command)
        self.assertIn('--headless', command)

    def test_each_conversion_has_an_independent_profile(self):
        profiles = []

        def inspect_profile(**kwargs):
            command = self.popen.call_args.args[0]
            profile_uri = next(arg.split('=', 1)[1] for arg in command
                               if arg.startswith('-env:UserInstallation='))
            profiles.append(profile_uri)
            output_dir = Path(command[command.index('--outdir') + 1])
            profile_dir = output_dir / 'profile'
            self.assertEqual(profile_dir.as_uri(), profile_uri)
            self.assertIn('ExportHiddenSlides',
                          (profile_dir / 'user' / 'registrymodifications.xcu').read_text(encoding='utf-8'))
            return self.write_converted_pdf(**kwargs)

        self.process.communicate.side_effect = inspect_profile
        self.convert()
        self.convert()
        self.assertNotEqual(*profiles)
        self.assertEqual([], list(self.root.glob('reports-libreoffice-*')))

    def test_empty_or_invalid_pptx_never_starts_process(self):
        from pptx.exc import PackageNotFoundError

        for content in (pptx_bytes(count=0), b'not a pptx'):
            with self.subTest(content_length=len(content)):
                self.source.write_bytes(content)
                with self.assertRaises((ValueError, PackageNotFoundError)):
                    self.convert()
                self.assertFalse(self.target.exists())
                self.assertEqual(content, self.source.read_bytes())
        self.popen.assert_not_called()

    def test_nonzero_exit_removes_temporary_output(self):
        self.process.returncode = 1
        with self.assertRaisesRegex(RuntimeError, '退出码 1'):
            self.convert()
        self.assert_clean_failure()

    def test_zero_exit_without_pdf_is_failure(self):
        self.process.communicate.side_effect = None
        with self.assertRaisesRegex(ValueError, '未生成 PDF'):
            self.convert()
        self.assert_clean_failure()

    def test_invalid_encrypted_or_truncated_pdf_is_failure(self):
        with fitz.open(stream=pdf_bytes(), filetype='pdf') as document:
            encrypted = document.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256,
                                         owner_pw='test', user_pw='test')
            document.delete_page(1)
            truncated = document.tobytes()
        for output in (b'invalid pdf', encrypted, truncated):
            with self.subTest(output_length=len(output)):
                self.output = output
                with self.assertRaises((ValueError, fitz.FileDataError)):
                    self.convert()
                self.assert_clean_failure()

    def test_failed_conversion_does_not_overwrite_existing_pdf(self):
        self.target.write_bytes(b'existing pdf')
        self.process.returncode = 1
        with self.assertRaises(RuntimeError):
            self.convert()
        self.assertEqual(b'existing pdf', self.target.read_bytes())

    def test_process_start_failure_cleans_profile_and_preserves_source(self):
        self.popen.side_effect = OSError('cannot start')
        with self.assertRaises(OSError):
            self.convert()
        self.assert_clean_failure()

    def test_timeout_terminates_only_this_conversion_process_tree(self):
        for platform in ('win32', 'linux'):
            with self.subTest(platform=platform), \
                    patch.dict(self.namespace, {'sys': SimpleNamespace(platform=platform)}), \
                    patch('subprocess.CREATE_NO_WINDOW', 0x08000000, create=True), \
                    patch('signal.SIGKILL', 9, create=True), \
                    patch('subprocess.run') as taskkill, \
                    patch('os.killpg', create=True) as killpg:
                self.process.communicate.side_effect = [
                    subprocess.TimeoutExpired('test-soffice', 120), (b'', b'')]
                with self.assertRaisesRegex(RuntimeError, '转换超时'):
                    self.convert()
                if platform == 'win32':
                    self.assertEqual(['taskkill', '/PID', '43210', '/T', '/F'],
                                     taskkill.call_args.args[0])
                    self.assertEqual(subprocess.CREATE_NO_WINDOW, self.popen.call_args.kwargs['creationflags'])
                    killpg.assert_not_called()
                else:
                    killpg.assert_called_once_with(43210, signal.SIGKILL)
                    self.assertTrue(self.popen.call_args.kwargs['start_new_session'])
                    taskkill.assert_not_called()
                self.assert_clean_failure()


if __name__ == '__main__':
    unittest.main()
