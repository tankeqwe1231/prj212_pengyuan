# -*- coding: utf-8 -*-
from __future__ import absolute_import
from __future__ import division
from __future__ import print_function
from __future__ import unicode_literals
import time
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options
from PIL import Image
import platform
import io
import os
from datetime import datetime
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_VERTICAL_ANCHOR
import logging


def remove_white_bg(img):
    datas = img.getdata()
    new_data = []
    for item in datas:
        # 判断是否为接近白色的像素，可根据实际情况调整阈值
        if item[0] > 240 and item[1] > 240 and item[2] > 240:
            new_data.append((255, 255, 255, 0))  # 变为透明
        else:
            new_data.append(item)
    img.putdata(new_data)
    return img


class ScreenShotToPPTX:
    def __init__(self):
        self.output_folder = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'assets', 'docs')
        self.driver = None
        self.slide_list = []
        self.ppt_template_dir = os.path.join(os.path.dirname(__file__), 'ppt_template')
        self.system = platform.system()
        # 创建输出文件夹
        if not os.path.exists(self.output_folder):
            os.makedirs(self.output_folder)
        self.setup_driver()

    def setup_driver(self):
        """设置并启动Chrome浏览器"""
        chrome_options = Options()
        chrome_options.add_argument('--headless')
        chrome_options.add_argument('--hide-scrollbars')
        # headless模式下--start-maximized无效，明确设置窗口大小
        chrome_options.add_argument('--window-size=1920,1080')
        if self.system == 'Linux':
            chrome_options.add_argument("--no-sandbox")
        self.driver = webdriver.Chrome(options=chrome_options)

    def load_page(self, url):
        """加载页面并设置窗口大小"""
        logging.info(f"加载页面: {url}", datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
        self.driver.get(url)
        # 等待所有AJAX请求完成 - 基于react-grid-item元素数量 和 performance API
        timeout = 60  # 最多等待60秒
        start_time = time.time()
        grid_items_count = 0
        completed_ajax_count = 0
        logging.info("开始ajax请求完成情况监控 ", datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
        while time.time() - start_time < timeout:
            try:
                # 获取当前grid-item元素数量
                grid_items_count = len(self.driver.find_elements(By.CSS_SELECTOR, ".react-grid-item"))
                # 使用JavaScript获取已完成的XHR请求
                get_xhr_script = """
                    return window.performance.getEntriesByType('resource')
                        .filter(entry => entry.initiatorType === 'xmlhttprequest').length;
                """
                completed_ajax_count = self.driver.execute_script(get_xhr_script)
                print(f"当前已完成AJAX请求数: {completed_ajax_count}, grid-item元素数: {grid_items_count}")
                # 如果AJAX请求数大于等于元素数，认为页面已完全加载
                if completed_ajax_count >= grid_items_count > 0:
                    logging.info("AJAX请求数量已满足条件，页面加载完成")
                    break
                # 短暂等待后再次检查
                time.sleep(1)
            except Exception as e:
                logging.error(f"监控网络请求时出错: {e}")
                # 如果发生错误，可能需要不同的等待策略或增加等待时间
                time.sleep(2)  # 发生错误时等待更久一点

        logging.info(f"网络监控结束，已完成AJAX请求数: {completed_ajax_count}，grid-item元素数: {grid_items_count}",
              datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
        # 计算页面尺寸
        page_width = self.driver.execute_script('return document.documentElement.scrollWidth')
        page_height = self.driver.execute_script(
            "return Math.max(document.body.scrollHeight, document.documentElement.clientHeight, document.documentElement.scrollHeight, document.documentElement.offsetHeight);")
        page_height = page_height + 20  # 增加额外高度以防万一
        self.driver.set_window_size(page_width, page_height)
        logging.info(f"设置窗口大小为: {page_width}x{page_height}")
        time.sleep(2)  # 等待窗口调整生效

    def capture_grouped_screenshots(self, url, dash_id):
        self.load_page(url)
        """按照markup_element分组截取图片"""
        # 查找所有markup_element元素
        markup_elements = self.driver.find_elements(By.CSS_SELECTOR, ".react-grid-item.markup")
        if len(markup_elements) == 0:
            raise Exception("未找到markup元素")
        # 查找所有react-grid-item元素
        all_grid_items = self.driver.find_elements(By.CSS_SELECTOR, ".react-grid-item")

        # 添加最后一个虚拟分隔符，以便处理最后一组元素
        markup_indices = [all_grid_items.index(element) for element in markup_elements]
        markup_indices.append(len(all_grid_items))  # 添加末尾索引

        # 对每组元素进行截图
        for i in range(len(markup_indices) - 1):
            # 当前分隔符的索引
            current_separator_idx = markup_indices[i]
            # 下一个分隔符的索引
            next_separator_idx = markup_indices[i + 1]
            # 如果下一个分隔符紧跟着当前分隔符，则没有内容元素
            if next_separator_idx - current_separator_idx <= 1:
                continue
            slide_info = {"title": "", "img": ""}
            markup_element = all_grid_items[current_separator_idx]

            # 获取标题值和iframe内容
            # 获取editable-title中input的value值
            title_input = markup_element.find_element(By.CSS_SELECTOR, ".editable-title input")
            title_value = title_input.get_attribute("value")
            slide_info["title"] = title_value

            # # 获取iframe内的内容
            # iframe = markup_element.find_element(By.CSS_SELECTOR, "iframe.cke_wysiwyg_frame")
            # # 切换到iframe
            # self.driver.switch_to.frame(iframe)
            # # 获取body内的HTML
            # body_html = self.driver.find_element(By.TAG_NAME, "body").get_attribute("innerHTML")
            # slide_info["text"] = body_html
            # # 切回主文档
            # self.driver.switch_to.default_content()
            # print(f"已获取分隔元素的标题: {title_value} 和内容文本")

            # 获取当前组内的所有内容元素
            group_elements = all_grid_items[current_separator_idx + 1:next_separator_idx]
            # 如果组内没有元素，继续下一组
            if not group_elements:
                continue
            # 获取组内元素的ID
            group_ids = [element.get_attribute('id') for element in group_elements]
            logging.info(f"开始截屏分组元素: {group_ids}")

            # 计算包含所有元素的边界框
            min_x = float('inf')
            min_y = float('inf')
            max_right = 0
            max_bottom = 0

            # 对组内的每个元素获取位置信息
            for element in group_elements:
                rect = self.driver.execute_script("""
                    var rect = arguments[0].getBoundingClientRect();
                    return {
                        'x': rect.left,
                        'y': rect.top,
                        'width': rect.width,
                        'height': rect.height
                    };
                """, element)

                min_x = min(min_x, rect['x'])
                min_y = min(min_y, rect['y'])
                max_right = max(max_right, rect['x'] + rect['width'])
                max_bottom = max(max_bottom, rect['y'] + rect['height'])

            # 获取整个元素组的屏幕截图
            full_screenshot = self.capture_full_screenshot(min_x, min_y, max_right, max_bottom)
            slide_info['img'] = full_screenshot
            slide_info['img_count'] = len(group_elements)
            self.slide_list.append(slide_info)

            # 生成文件名和保存路径
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            # 使用第一个元素的ID作为组ID
            group_id = group_ids[0].replace('slice_', '')
            # 在文件名中加入看板id，避免不同页面同名文件覆盖
            filename = f"dash{dash_id}_slice{group_id}_{timestamp}.png"
            save_path = os.path.join(self.output_folder, filename)
            # full_screenshot.save(save_path)
            # print(f"已保存分组截图: {save_path}")

    def capture_full_screenshot(self, min_x, min_y, max_right, max_bottom):
        """直接捕获整个元素区域的截图，不需要滚动和拼接"""
        # 获取设备像素比例
        pixel_ratio = self.driver.execute_script('return window.devicePixelRatio;')
        # 获取屏幕截图
        screenshot = self.driver.get_screenshot_as_png()
        image = Image.open(io.BytesIO(screenshot))
        # 计算裁剪区域，考虑设备像素比
        crop_left = min_x * pixel_ratio
        crop_top = min_y * pixel_ratio
        crop_right = max_right * pixel_ratio
        crop_bottom = max_bottom * pixel_ratio
        # 确保裁剪区域有效
        if crop_left < 0:
            crop_left = 0
        if crop_top < 0:
            crop_top = 0
        if crop_right > image.width:
            crop_right = image.width
        if crop_bottom > image.height:
            crop_bottom = image.height
        # 裁剪并返回图片
        return image.crop((crop_left, crop_top, crop_right, crop_bottom))

    def save_to_ppt(self, task_info, ppt_path):
        """将PNG图像列表保存到PowerPoint文件中"""
        logging.info(f"save_to_ppt 开始创建PPT: {ppt_path}")

        # 模版ppt路径
        ppt_template_path = os.path.join(self.ppt_template_dir, task_info['task_ppt_template_name'])
        # 从模版创建Presentation对象
        prs = Presentation(ppt_template_path)
        # 获取PowerPoint布局
        slide_layouts = prs.slide_layouts

        # home Slide 布局
        home_slide_layout = slide_layouts[0]
        # 内容页布局占位符信息
        slide_layout_content = slide_layouts[1]
        # 尾页slide布局
        slide_layout_end = slide_layouts[len(slide_layouts) - 1]

        home_slide_layout_placeholder_info = {}
        slide_layout_content_placeholder_info = {}
        print("首页布局占位符信息:")
        for shape in home_slide_layout.placeholders:
            print(f"占位符idx={shape.placeholder_format.idx}, type={shape.placeholder_format.type}, name={shape.name}")
            if shape.has_text_frame:
                text = shape.text_frame.text.strip().lower()
                print(f"text_frame={text}")
                if text:
                    home_slide_layout_placeholder_info[shape.placeholder_format.idx] = text
        print("内容页布局占位符信息:")
        for shape in slide_layout_content.placeholders:
            print(f"占位符idx={shape.placeholder_format.idx}, type={shape.placeholder_format.type}, name={shape.name}")
            if shape.has_text_frame:
                text = shape.text_frame.text.strip().lower()
                print(f"text_frame={text}")
                if text:
                    slide_layout_content_placeholder_info[shape.placeholder_format.idx] = text
                    if text == "title":
                        slide_layout_content_placeholder_info[text] = {"top": shape.top, "height": shape.height}

        slide_width = prs.slide_width
        slide_height = prs.slide_height
        # 添加首张幻灯片,加上背景图
        slide_home = prs.slides.add_slide(home_slide_layout)
        # 通过idx匹配母版中的占位符并应用
        for shape in slide_home.placeholders:
            if shape.has_text_frame:
                idx = shape.placeholder_format.idx
                print(f"处理首页内容占位符: idx={idx}, name={shape.name}")
                if idx in home_slide_layout_placeholder_info:
                    if home_slide_layout_placeholder_info[idx] == "title":
                        shape.text = task_info['task_ppt_home_title']
                    elif home_slide_layout_placeholder_info[idx] == "date":
                        shape.text = datetime.now().strftime('%Y/%m/%d')

        dpi = 96
        # 标题上方留白
        top_title = Inches(6 / dpi)
        # 标题下方留白，暂时不用
        # bottom_title = Inches(15/dpi)
        # 标题高度
        title_height = Inches(68 / dpi)
        # top_height 内容slide顶部已占用高度
        top_height = top_title + title_height
        if 'title' in slide_layout_content_placeholder_info:
            title_position = slide_layout_content_placeholder_info['title']
            top_height = title_position['top'] + title_position['height']
        # slide 底部留白
        bottom_height = Inches(15 / dpi)

        # 计算实际可用于放置图片的空间尺寸
        available_width = slide_width
        available_height = slide_height - top_height - bottom_height
        # 计算可用空间的宽高比
        slide_ratio = available_width / available_height
        logging.info(f"内容slide可用空间比例: 宽={available_width}, 高={available_height}, 比例={slide_ratio}")

        # 遍历图像列表
        for i, slide_dict in enumerate(self.slide_list):
            print(f"开始添加slide {i + 1}/{len(self.slide_list)}")
            # 添加新的带母版布局的幻灯片
            slide = prs.slides.add_slide(slide_layout_content)
            # 通过idx匹配内容页中的占位符并应用
            for shape in slide.placeholders:
                if shape.has_text_frame:
                    idx = shape.placeholder_format.idx
                    if idx in slide_layout_content_placeholder_info:
                        if slide_layout_content_placeholder_info[idx] == "title":
                            shape.text = slide_dict['title']
                            # 设置文本垂直居中
                            # shape.text_frame.vertical_anchor = MSO_VERTICAL_ANCHOR.CENTER

            img = slide_dict['img']
            if img:
                # 如果需要去掉背景,需要转换图像模式为RGBA以支持透明度
                if img.mode != 'RGBA':
                    img = img.convert('RGBA')
                img = remove_white_bg(img)
                # 创建临时文件以保存图像
                temp_image_path = os.path.join(self.output_folder, f"temp_image_{task_info['taskid']}_{i}.png")
                img.save(temp_image_path, format='PNG')
                # img.save(temp_image_path)
                # 计算图像宽高比例以适应幻灯片
                img_width, img_height = img.size
                img_ratio = img_width / img_height
                print(f"图片宽高比: 宽={img_width}, 高={img_height}, 比例={img_ratio}")
                # 计算图像在幻灯片上的位置和大小，按比例重新计算
                # 实际图片宽高比大于幻灯片宽高比，根据slide实际宽度重新计算图片高度pic_height
                if img_ratio > slide_ratio:
                    # 在可用宽度上预留出一点空间
                    pic_width = available_width - Inches(0.2)
                    pic_height = pic_width / img_ratio
                    left = (slide_width - pic_width) / 2
                    top = top_height + (available_height - pic_height) / 2
                    # top = top_height
                else:
                    pic_height = available_height - Inches(0.2)
                    pic_width = pic_height * img_ratio
                    # top = top_height
                    top = top_height + (available_height - pic_height) / 2
                    left = (slide_width - pic_width) / 2
                # 添加图片到幻灯片
                slide.shapes.add_picture(temp_image_path, left, top, width=pic_width, height=pic_height)
                # 删除临时图像文件
                os.remove(temp_image_path)
            print(f"已添加slide {i + 1}/{len(self.slide_list)}")

        # 如果slide_layouts长度>3,则从下标2开始（包含）至末尾（不包含）的布局直接添加
        if len(slide_layouts) > 3:
            for i in range(2, len(slide_layouts) - 1):
                prs.slides.add_slide(slide_layouts[i])
        # 添加尾页幻灯片
        prs.slides.add_slide(slide_layout_end)

        # 保存PowerPoint文件
        prs.save(ppt_path)
        logging.info(f"ppt创建成功: {ppt_path}")
